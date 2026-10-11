"""PowerPoint parser (.pptx and .ppt).

Extracts per-slide content from presentations:
- Slide title (from title placeholder or first text shape)
- All text from text frames
- Tables (rendered as structured text rows)
- Speaker notes

.ppt (legacy binary) is not natively supported by python-pptx.
The parser detects this and raises a clear error rather than silently failing.
"""

from __future__ import annotations

from pathlib import Path

from rag_chatbot.config.constants import PPTX_EXTENSIONS
from rag_chatbot.utils.exception_handler import DocumentParseError, InvalidDocumentContentError, UnsupportedFormatError
from rag_chatbot.utils.logger import logger

from .base_parser import BaseParser, ParsedChunk


class PptxParser(BaseParser):
    """Parses PowerPoint presentations into per-slide ParsedChunk objects."""

    @property
    def supported_extensions(self) -> frozenset[str]:
        return PPTX_EXTENSIONS

    def parse(self, file_path: Path, filename: str) -> list[ParsedChunk]:
        ext = file_path.suffix.lstrip(".").lower()
        logger.info("Parsing PowerPoint file: %s (ext=%s)", filename, ext)

        if ext == "ppt":
            raise UnsupportedFormatError(
                f"Legacy .ppt format is not directly supported. "
                f"Please convert {filename!r} to .pptx using Microsoft PowerPoint or LibreOffice "
                f"before uploading.",
                details={"filename": filename, "ext": ext},
            )

        try:
            return self._parse_pptx(file_path, filename)
        except (DocumentParseError, InvalidDocumentContentError, UnsupportedFormatError):
            raise
        except Exception as exc:
            raise DocumentParseError(
                f"Failed to parse PowerPoint file {filename!r}: {exc}",
                details={"filename": filename},
            ) from exc

    def _parse_pptx(self, file_path: Path, filename: str) -> list[ParsedChunk]:
        try:
            from pptx import Presentation  # noqa: PLC0415
            from pptx.util import Pt  # noqa: PLC0415, F401
        except ImportError as exc:
            raise DocumentParseError("python-pptx is required for .pptx parsing.") from exc

        try:
            prs = Presentation(str(file_path))
        except Exception as exc:
            raise DocumentParseError(f"Cannot open .pptx file {filename!r}: {exc}") from exc

        chunks: list[ParsedChunk] = []

        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_title = self._extract_title(slide)
            slide_text_parts: list[str] = []

            if slide_title:
                slide_text_parts.append(f"Slide Title: {slide_title}")

            # Text from all text frames (excluding title shape already captured)
            for shape in slide.shapes:
                if shape.has_text_frame:
                    shape_text = self._extract_text_frame(shape.text_frame)
                    if shape_text and shape_text != slide_title:
                        slide_text_parts.append(shape_text)

                # Tables
                if shape.has_table:
                    table_text = self._extract_table(shape.table)
                    if table_text:
                        slide_text_parts.append(table_text)

            # Speaker notes
            notes_text = self._extract_notes(slide)
            if notes_text:
                slide_text_parts.append(f"Notes: {notes_text}")

            full_content = "\n\n".join(p for p in slide_text_parts if p.strip())

            if not full_content.strip():
                # Skip blank slides silently
                logger.debug("Slide %d in %r is blank, skipping", slide_idx, filename)
                continue

            chunks.append(
                ParsedChunk(
                    content=full_content,
                    source_filename=filename,
                    file_type="pptx",
                    slide_number=slide_idx,
                    slide_title=slide_title,
                )
            )

        if not chunks:
            raise InvalidDocumentContentError(f"PowerPoint file {filename!r} contains no slides with extractable text.")

        logger.debug("PPTX %r: %d slide chunks", filename, len(chunks))
        return chunks

    # ------------------------------------------------------------------
    # Extraction helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _extract_title(slide: object) -> str | None:
        """Extract the slide title from the title placeholder if present."""
        try:
            for ph in slide.placeholders:  # type: ignore[attr-defined]
                if ph.placeholder_format.idx == 0:  # 0 = title
                    text = ph.text.strip()
                    return text if text else None
        except Exception as exc:  # noqa: BLE001
            logger.debug("Could not extract title from placeholder: %s", exc)
        # Fallback: first shape with text
        try:
            for shape in slide.shapes:  # type: ignore[attr-defined]
                if shape.has_text_frame:
                    text = shape.text.strip()
                    if text:
                        return text
        except Exception as exc:  # noqa: BLE001
            logger.debug("Could not extract title from shapes: %s", exc)
        return None

    @staticmethod
    def _extract_text_frame(text_frame: object) -> str:
        """Extract all paragraph text from a text frame."""
        lines: list[str] = []
        try:
            for para in text_frame.paragraphs:  # type: ignore[attr-defined]
                line = para.text.strip()
                if line:
                    lines.append(line)
        except Exception as exc:  # noqa: BLE001
            logger.debug("Could not extract text frame: %s", exc)
        return "\n".join(lines)

    @staticmethod
    def _extract_table(table: object) -> str:
        """Render a table as pipe-delimited text."""
        rows: list[str] = []
        try:
            for row in table.rows:  # type: ignore[attr-defined]
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(" | ".join(cells))
        except Exception as exc:  # noqa: BLE001
            logger.debug("Could not extract table: %s", exc)
        if rows:
            return "Table:\n" + "\n".join(rows)
        return ""

    @staticmethod
    def _extract_notes(slide: object) -> str:
        """Extract speaker notes from a slide, if any."""
        try:
            notes_slide = slide.notes_slide  # type: ignore[attr-defined]
            if notes_slide and notes_slide.notes_text_frame:
                text = notes_slide.notes_text_frame.text.strip()
                return text if text else ""
        except Exception as exc:  # noqa: BLE001
            logger.debug("Could not extract speaker notes: %s", exc)
        return ""
